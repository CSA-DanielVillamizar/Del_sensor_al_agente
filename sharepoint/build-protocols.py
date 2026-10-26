"""T-028 HU-07 CA-07.3: convierte los protocolos de Markdown a .docx para la biblioteca Protocolos.

El conocimiento de SharePoint de Copilot Studio indexa .docx y no Markdown (plan §4.4).
El conversor admite solo el subconjunto de Markdown que usan los protocolos:

- `# Título`            -> estilo Title
- `## Apartado`         -> estilo Heading 1
- `1. Paso`             -> estilo List Number (lista numerada real de Word)
- `- Punto`             -> estilo List Bullet
- `**negrita**`         -> texto en negrita dentro de un párrafo
- cualquier otra línea  -> párrafo normal

Uso:
    python sharepoint/build-protocols.py [--src protocols] [--out protocols/dist]
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = ROOT / "protocols"
DEFAULT_OUT = ROOT / "protocols" / "dist"

NUMBERED = re.compile(r"^\d+\.\s+(.*)$")
BULLET = re.compile(r"^[-*]\s+(.*)$")
BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_inline(paragraph, text):
    """Agrega el texto al párrafo, con las partes entre ** en negrita."""
    position = 0
    for match in BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _new_document():
    """Plantilla fija: A4, márgenes de 2 cm y letra de 11 puntos, para que quepa en una página."""
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    doc.styles["Normal"].font.size = Pt(11)
    return doc


def convert(md_path: Path, out_path: Path) -> Path:
    """Convierte un protocolo Markdown en un .docx con encabezados y listas reales."""
    doc = _new_document()
    title = None
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
        elif line.startswith("# "):
            title = line[2:].strip()
            doc.add_paragraph(title, style="Title")
        elif match := NUMBERED.match(line):
            _add_inline(doc.add_paragraph(style="List Number"), match.group(1))
        elif match := BULLET.match(line):
            _add_inline(doc.add_paragraph(style="List Bullet"), match.group(1))
        else:
            _add_inline(doc.add_paragraph(), line)

    doc.core_properties.title = title or md_path.stem
    doc.core_properties.subject = "Protocolo ficticio para la demo"
    doc.core_properties.author = "Del sensor al agente"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build(src_dir: Path = DEFAULT_SRC, out_dir: Path = DEFAULT_OUT) -> list[Path]:
    """Convierte cada protocolo de `src_dir` y devuelve las rutas de los .docx generados."""
    sources = sorted(Path(src_dir).glob("*.md"))
    if not sources:
        raise FileNotFoundError(f"No hay protocolos .md en {src_dir}")
    return [convert(md, Path(out_dir) / f"{md.stem}.docx") for md in sources]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Convierte los protocolos Markdown a .docx.")
    parser.add_argument("--src", type=Path, default=DEFAULT_SRC, help="carpeta con los .md")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="carpeta de salida de los .docx")
    args = parser.parse_args(argv)
    try:
        outputs = build(args.src, args.out)
    except FileNotFoundError as error:
        print(error, file=sys.stderr)
        return 1
    for path in outputs:
        print(f"Generado: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
